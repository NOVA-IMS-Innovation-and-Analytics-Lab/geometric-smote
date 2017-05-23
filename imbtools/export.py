import docx

def export_table(df, docx_path):
    doc = docx.Document(docx_path)
    table = doc.add_table(df.shape[0]+1, df.shape[1])
    for j in range(df.shape[1]):
        table.cell(0, j).text = df.columns[j]
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            table.cell(i+1, j).text = str(df.values[i,j])
    doc.save(docx_path)